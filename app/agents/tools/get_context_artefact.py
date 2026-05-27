"""``get_context_artefact`` — file parser for the extractor agent.

M2 shipped the extractor as a scaffold with this tool listed but not
implemented. M5 implements the parser backend: takes an S3 object key
(or raw bytes for tests) and returns a ``ContextArtefact`` with the
text the LLM should read PLUS base64-encoded images for Claude's
vision API.

Supported kinds:
- ``pdf`` — pypdf page-by-page text extraction.
- ``docx`` — python-docx paragraphs + table cells.
- ``image`` — PNG/JPEG inlined as base64 for vision (no OCR).
- ``text`` — plain UTF-8.

Anything outside these gets ``kind="unknown"`` with empty text — the
extractor will refuse to invent fields when source text is missing
(per its system prompt).

S3 fetch uses the shared boto3 client from ``app/utils/s3.py``; the
sync call runs in a thread via ``asyncio.to_thread`` so the FastAPI
event loop stays responsive.
"""

from __future__ import annotations

import asyncio
import base64
import io
from dataclasses import dataclass, field
from typing import Any, Literal

import docx
from pypdf import PdfReader

from app.config import settings
from app.errors import BusinessRuleError, ValidationError
from app.utils.logging import get_logger
from app.utils.s3 import get_s3_client

log = get_logger(__name__)

# Per Claude-vision docs: 5 MB hard cap per image; 20 MB total per request.
_MAX_IMAGE_BYTES = 5 * 1024 * 1024
_MAX_TOTAL_IMAGE_BYTES = 20 * 1024 * 1024

_ALLOWED_IMAGE_MIMES: frozenset[str] = frozenset({"image/png", "image/jpeg", "image/jpg"})

ContextKind = Literal["pdf", "docx", "image", "text", "unknown"]


@dataclass(frozen=True)
class InlineImage:
    """One image attachment for Claude vision (base64-encoded)."""

    media_type: str
    data_b64: str
    size_bytes: int


@dataclass(frozen=True)
class ContextArtefact:
    """Result of parsing one uploaded file.

    The extractor agent receives this via its ``get_context_artefact``
    tool. ``text`` is the primary input; ``attachments`` go into the
    multimodal content block when the LLM is called.
    """

    artefact_id: str
    kind: ContextKind
    text: str
    attachments: list[InlineImage] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)


def _sniff_mime(filename: str | None, content_type: str | None) -> str:
    """Best-effort MIME resolution from explicit content type, then extension."""
    if content_type:
        return content_type.lower()
    if filename:
        lower = filename.lower()
        if lower.endswith(".pdf"):
            return "application/pdf"
        if lower.endswith(".docx"):
            return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        if lower.endswith(".png"):
            return "image/png"
        if lower.endswith((".jpg", ".jpeg")):
            return "image/jpeg"
        if lower.endswith(".txt"):
            return "text/plain"
    return "application/octet-stream"


def parse_pdf_bytes(data: bytes) -> tuple[str, int]:
    """Extract plain text from a PDF byte blob. Returns ``(text, page_count)``.

    pypdf can be memory-heavy on large PDFs; the caller MUST enforce a
    size cap before calling (see ``settings.upload_max_file_size_mb``).
    """
    reader = PdfReader(io.BytesIO(data))
    chunks: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            chunks.append(page_text)
    return "\n\n".join(chunks), len(reader.pages)


def parse_docx_bytes(data: bytes) -> str:
    """Extract paragraph + table text from a DOCX byte blob."""
    document = docx.Document(io.BytesIO(data))
    chunks: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            chunks.append(para.text)
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                chunks.append(row_text)
    return "\n\n".join(chunks)


def image_to_inline(data: bytes, mime: str) -> InlineImage:
    """Base64-encode an image for Claude vision."""
    if mime not in _ALLOWED_IMAGE_MIMES:
        raise ValidationError(
            f"image MIME {mime!r} not supported (allowed: {sorted(_ALLOWED_IMAGE_MIMES)})",
            field="media_type",
        )
    if len(data) > _MAX_IMAGE_BYTES:
        raise ValidationError(
            f"image exceeds {_MAX_IMAGE_BYTES // (1024 * 1024)} MB per-image limit",
            field="size_bytes",
            detail={"received": len(data), "limit": _MAX_IMAGE_BYTES},
        )
    return InlineImage(
        media_type="image/jpeg" if mime == "image/jpg" else mime,
        data_b64=base64.standard_b64encode(data).decode("ascii"),
        size_bytes=len(data),
    )


def _fetch_s3_bytes(bucket: str, key: str) -> tuple[bytes, str]:
    """Sync S3 GET. Returns ``(body, content_type)``. Used via asyncio.to_thread."""
    client = get_s3_client()
    response = client.get_object(Bucket=bucket, Key=key)
    body: bytes = response["Body"].read()
    content_type: str = response.get("ContentType") or "application/octet-stream"
    return body, content_type


async def _fetch_s3_async(bucket: str, key: str) -> tuple[bytes, str]:
    """Run the blocking S3 GET on a worker thread."""
    return await asyncio.to_thread(_fetch_s3_bytes, bucket, key)


def _check_size_cap(data: bytes, kind: str) -> None:
    limit_mb = settings.upload_max_file_size_mb
    limit_bytes = limit_mb * 1024 * 1024
    if len(data) > limit_bytes:
        raise ValidationError(
            f"{kind} exceeds {limit_mb} MB upload cap",
            field="size_bytes",
            detail={"received": len(data), "limit_bytes": limit_bytes},
        )


async def get_context_artefact(
    artefact_id: str,
    *,
    s3_key: str | None = None,
    s3_bucket: str | None = None,
    raw_bytes: bytes | None = None,
    filename: str | None = None,
    content_type: str | None = None,
) -> ContextArtefact:
    """Fetch + parse an upload into a structured ``ContextArtefact``.

    Two input modes:
    - ``s3_key`` (and optionally ``s3_bucket``) — fetch from MinIO/S3.
    - ``raw_bytes`` (with optional ``filename`` + ``content_type``) —
      direct injection, used by tests and the CLI wizard.

    Exactly one of (s3_key, raw_bytes) must be supplied.
    """
    if (s3_key is None) == (raw_bytes is None):
        raise BusinessRuleError(
            "exactly one of s3_key or raw_bytes must be supplied",
            detail={"artefact_id": artefact_id},
        )

    if s3_key is not None:
        bucket = s3_bucket or settings.s3_bucket
        body, server_mime = await _fetch_s3_async(bucket, s3_key)
        mime = _sniff_mime(filename or s3_key, content_type or server_mime)
    else:
        assert raw_bytes is not None
        body = raw_bytes
        mime = _sniff_mime(filename, content_type)

    _check_size_cap(body, kind="upload")

    if mime == "application/pdf":
        text, page_count = parse_pdf_bytes(body)
        return ContextArtefact(
            artefact_id=artefact_id,
            kind="pdf",
            text=text,
            metadata={"mime": mime, "size_bytes": len(body), "page_count": page_count},
        )

    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        text = parse_docx_bytes(body)
        return ContextArtefact(
            artefact_id=artefact_id,
            kind="docx",
            text=text,
            metadata={"mime": mime, "size_bytes": len(body)},
        )

    if mime in _ALLOWED_IMAGE_MIMES:
        inline = image_to_inline(body, mime)
        return ContextArtefact(
            artefact_id=artefact_id,
            kind="image",
            text="",
            attachments=[inline],
            metadata={"mime": inline.media_type, "size_bytes": len(body)},
        )

    if mime.startswith("text/"):
        return ContextArtefact(
            artefact_id=artefact_id,
            kind="text",
            text=body.decode("utf-8", errors="replace"),
            metadata={"mime": mime, "size_bytes": len(body)},
        )

    log.warning("artefact_unknown_mime", artefact_id=artefact_id, mime=mime, size=len(body))
    return ContextArtefact(
        artefact_id=artefact_id,
        kind="unknown",
        text="",
        metadata={"mime": mime, "size_bytes": len(body)},
    )
